"""Read receipts / invoices / bank-statement screenshots with Gemini and return transaction rows.

Nothing is saved here: the rows go to a review screen and are only stored after the user confirms.
"""
import csv
import io
import os
import unicodedata
import zipfile
from datetime import date, datetime

from pydantic import BaseModel, Field

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_MODEL = "gemini-2.5-flash"
MAX_SIDE = 2200                      # downscale big phone photos: fewer tokens, same readability
MAX_BYTES = 10 * 1024 * 1024         # per file
MAX_PDF_PAGES = 40                   # a few months of a bank statement
CHUNK_PAGES = 5                      # pages per Gemini call: small chunks keep long statements accurate
PAGES_PER_SCAN = 10                  # quota: one scan per started 10 pages
MAX_PIXELS = 50_000_000              # refuse decompression bombs
MAX_TABLE_ROWS = 2000                # spreadsheet / CSV bank exports
TABLE_CHUNK_ROWS = 200               # rows per Gemini call, and per scan of quota
MAX_UNZIPPED = 50 * 1024 * 1024      # an .xlsx is a zip: refuse zip bombs
MAX_CELL = 200


class Row(BaseModel):
    tdate: str = Field(description="Transaction date as YYYY-MM-DD")
    amount: float = Field(description="Positive amount in euro, e.g. 12.5")
    kind: str = Field(description="'expense' for money out, 'income' for money in, 'transfer' for money moved "
                                  "between the user's own accounts/pockets (top-ups, savings, vaults)")
    description: str = Field(description="Short merchant / payee / description, max 60 chars")
    category: str = Field(description="Exactly one of the given category names for this kind, or '' if unsure")


class ScanResult(BaseModel):
    document_type: str = Field(description="'receipt', 'invoice', 'bank_statement' or 'other'")
    rows: list[Row]


PROMPT = """You read Greek (and English) financial documents for a personal budget app.
Today is {today}. The document is a photo, screenshot, PDF or spreadsheet export (given as tab-separated text) of ONE of:
- a receipt (απόδειξη) or invoice / bill (τιμολόγιο, λογαριασμός ρεύματος/νερού/κινητού)
  -> return exactly ONE row: the final amount paid (ΣΥΝΟΛΟ / ΠΛΗΡΩΤΕΟ / Total, VAT included),
     the document date, and the shop or company name as description.
- a bank statement, card statement, banking-app screenshot or bank export (Revolut, Πειραιώς, Εθνική, Eurobank, Alpha, TBI...)
  -> return ONE row PER transaction. Money out = 'expense', money in (salary, transfer in, refund) = 'income'.
     Skip balance lines, totals, the whole 'Pending' section (those appear again once completed), reversed holds,
     and currency-exchange duplicate lines.
     Money moved between the user's OWN accounts is kind 'transfer', never income or expense: card/Apple Pay/
     Google Pay top-ups ("Top-up by *1234"), "To pocket" / "From pocket", savings, vaults, round-ups, and the
     separate "Pockets" / "Savings" section of the statement (it repeats the same movements).
     From a normal bank account, card payments to Revolut or another e-money wallet (e.g. "Revolut **** Dublin")
     are top-ups of the user's own wallet: 'transfer'. Lines marked "Reallocated" or "TRANSFER TO/FROM ACCOUNT"
     between the user's own accounts are 'transfer'. "B/O <name>" (by order of) is money in from someone else.
- In spreadsheet exports a negative amount (or a Debit/Χρέωση column) is money out, positive (Credit/Πίστωση) is money in.

Rules:
- amount is always positive; the direction goes in kind.
- Greek numbers use comma for decimals (12,50 = 12.5) and dot for thousands (1.234,56 = 1234.56).
- Dates are usually DD/MM/YYYY. If the year is missing, use the most recent past date that fits (today is {today}).
- description: clean merchant name (e.g. "ΣΚΛΑΒΕΝΙΤΗΣ", "Lidl", "ΔΕΗ", "Wolt"), not the full raw bank text.
- category: copy one name EXACTLY (full text, including anything in parentheses) from these lists:
  expense categories: {expense_cats}
  income categories: {income_cats}
  Use '' if nothing fits. Hints: supermarket -> groceries category, cafe/coffee -> coffee category,
  fuel station (Shell, EKO, BP, Avin, Aegean) -> fuel, ΔΕΗ/ΕΥΔΑΠ/Cosmote/Vodafone/Nova -> bills,
  Netflix/Spotify/Google/Claude/Revolut plan -> subscriptions, pharmacy -> health.
- If the image is not a financial document or is unreadable, return document_type 'other' and no rows.
"""


def load_api_key():
    """GEMINI_API_KEY from the environment, else from app/.env."""
    key = os.environ.get("GEMINI_API_KEY", "").strip()
    if key:
        return key
    env_path = os.path.join(HERE, ".env")
    if os.path.exists(env_path):
        with open(env_path, encoding="utf-8") as f:
            for line in f:
                if line.strip().startswith("GEMINI_API_KEY="):
                    return line.split("=", 1)[1].strip().strip('"').strip("'")
    return ""


def sniff(data):
    """File type from the actual bytes (the browser's declared type and the file name are ignored)."""
    if data[:3] == bytes.fromhex("ffd8ff"):
        return "image/jpeg"
    if data[:8] == bytes.fromhex("89504e470d0a1a0a"):
        return "image/png"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    if data[:5] == b"%PDF-":
        return "application/pdf"
    if data[:4] == b"PK" + bytes.fromhex("0304"):
        return "application/zip"                      # possibly .xlsx; checked in prepare_upload
    if data[:8] == bytes.fromhex("d0cf11e0a1b11ae1"):
        return "application/x-ole"                    # old binary .xls
    if data[4:8] == b"ftyp" and data[8:12] in (b"heic", b"heix", b"mif1", b"msf1", b"heim", b"heis"):
        return "image/heic"
    return None


def prepare_upload(data, filename=""):
    """Validate an upload and make it safe to send out. Returns (bytes, mime); raises ValueError (Greek).

    Images are decoded and re-drawn onto a blank canvas, so EXIF (GPS location, camera serial) and any
    other embedded metadata or trailing payload are dropped. PDFs are parsed and page-limited.
    Spreadsheets (.xlsx) and CSV bank exports become plain tab-separated text (values only, no formulas)."""
    if not data:
        raise ValueError("Το αρχείο είναι άδειο")
    if len(data) > MAX_BYTES:
        raise ValueError("Το αρχείο ξεπερνά τα 10 MB")
    kind = sniff(data)
    if kind == "application/x-ole":
        raise ValueError("Τα παλιά αρχεία .xls δεν υποστηρίζονται — αποθήκευσέ το ως .xlsx ή CSV")
    if kind == "application/zip":
        return _table_text(_xlsx_rows(data)), "text/csv"
    if kind is None and filename.lower().endswith((".csv", ".tsv", ".txt")):
        return _table_text(_csv_rows(data)), "text/csv"
    if kind == "image/heic":
        raise ValueError("Οι φωτογραφίες HEIC δεν υποστηρίζονται — στείλε JPG ή PNG")
    if kind is None:
        raise ValueError("Το αρχείο δεν είναι εικόνα (JPG/PNG/WEBP), PDF, Excel (.xlsx) ή CSV")
    if kind == "application/pdf":
        from pypdf import PdfReader
        from pypdf.errors import PdfReadError
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("Το PDF είναι κλειδωμένο με κωδικό")
            pages = len(reader.pages)
        except (PdfReadError, OSError, KeyError, TypeError, AttributeError):
            raise ValueError("Το αρχείο δεν είναι έγκυρο PDF")
        if pages > MAX_PDF_PAGES:
            raise ValueError(f"Το PDF έχει {pages} σελίδες — το όριο είναι {MAX_PDF_PAGES}")
        return data, kind
    from PIL import Image, ImageOps
    Image.MAX_IMAGE_PIXELS = MAX_PIXELS
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
        img = ImageOps.exif_transpose(img)          # keep the photo upright before the EXIF is dropped
        img = img.convert("RGB")
    except (Image.DecompressionBombError, OSError, SyntaxError, ValueError):
        raise ValueError("Το αρχείο δεν είναι έγκυρη εικόνα")
    img.thumbnail((MAX_SIDE, MAX_SIDE))
    clean = Image.new("RGB", img.size, "white")
    clean.paste(img)
    out = io.BytesIO()
    clean.save(out, "JPEG", quality=88)
    return out.getvalue(), "image/jpeg"


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        v = v.date() if not (v.hour or v.minute or v.second) else v
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())[:MAX_CELL]


def _xlsx_rows(data):
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
        infos = z.infolist()
    except zipfile.BadZipFile:
        raise ValueError("Το αρχείο δεν είναι έγκυρο Excel (.xlsx)")
    if "xl/workbook.xml" not in z.namelist():
        raise ValueError("Το αρχείο δεν είναι εικόνα (JPG/PNG/WEBP), PDF, Excel (.xlsx) ή CSV")
    if len(infos) > 1000 or sum(i.file_size for i in infos) > MAX_UNZIPPED:
        raise ValueError("Το Excel είναι πολύ μεγάλο")
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)   # values only, never formulas
        best = []
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [_cell(v) for v in row[:30]]
                if any(cells):
                    rows.append(cells)
                if len(rows) > MAX_TABLE_ROWS + 60:
                    break
            if len(rows) > len(best):
                best = rows
        wb.close()
    except Exception:
        raise ValueError("Το αρχείο δεν είναι έγκυρο Excel (.xlsx)")
    return best


def _csv_rows(data):
    for enc in ("utf-8-sig", "cp1253"):                  # Greek banks often export in Windows-1253
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError("Το CSV δεν διαβάζεται ως κείμενο")
    controls = sum(1 for ch in text if not (ch.isprintable() or ch in "\r\n\t"))
    if chr(0) in text or controls > len(text) // 50:
        raise ValueError("Το CSV δεν διαβάζεται ως κείμενο")
    try:
        delim = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
    except csv.Error:
        delim = ";"
    rows = []
    for row in csv.reader(io.StringIO(text), delimiter=delim):
        cells = [_cell(v) for v in row[:30]]
        if any(cells):
            rows.append(cells)
        if len(rows) > MAX_TABLE_ROWS + 60:
            break
    return rows


def _header_index(rows):
    """Bank exports start with a few title lines; the table header is the first row with 3+ filled cells."""
    for i, r in enumerate(rows[:40]):
        if sum(1 for c in r if c) >= 3:
            return i
    return -1


def _table_text(rows):
    h = _header_index(rows)
    if not rows or len(rows) - (h + 1) <= 0:
        raise ValueError("Δεν βρέθηκαν γραμμές κινήσεων στο αρχείο")
    if len(rows) - (h + 1) > MAX_TABLE_ROWS:
        raise ValueError(f"Το αρχείο έχει πάνω από {MAX_TABLE_ROWS} γραμμές — χώρισέ το σε μικρότερα")
    return "\n".join("\t".join(r).rstrip("\t") for r in rows).encode("utf-8")


def split_table(data):
    """Tab-separated text -> chunks of TABLE_CHUNK_ROWS rows, each starting with the title lines + header."""
    lines = data.decode("utf-8").split("\n")
    h = _header_index([ln.split("\t") for ln in lines])
    head, body = lines[:h + 1], lines[h + 1:]
    return ["\n".join(head + body[i:i + TABLE_CHUNK_ROWS]).encode("utf-8")
            for i in range(0, len(body), TABLE_CHUNK_ROWS)] or [data]


def table_rows(parts):
    """Transaction rows across the chunks returned by split_table()."""
    total = 0
    for p in parts:
        lines = p.decode("utf-8").split("\n")
        total += len(lines) - (_header_index([ln.split("\t") for ln in lines]) + 1)
    return total


def scan_cost_rows(rows):
    """Spreadsheets: one scan per started TABLE_CHUNK_ROWS rows."""
    return max(1, -(-rows // TABLE_CHUNK_ROWS))


def split_pdf(data):
    """A (validated) PDF -> list of PDFs of at most CHUNK_PAGES pages each. Re-writing the pages also drops
    document-level extras (JavaScript, attachments, metadata) before anything leaves the server."""
    from pypdf import PdfReader, PdfWriter
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for start in range(0, len(reader.pages), CHUNK_PAGES):
        w = PdfWriter()
        for page in reader.pages[start:start + CHUNK_PAGES]:
            w.add_page(page)
        out = io.BytesIO()
        w.write(out)
        parts.append(out.getvalue())
    return parts


def pdf_pages(parts):
    """Total pages across the chunks returned by split_pdf()."""
    from pypdf import PdfReader
    return sum(len(PdfReader(io.BytesIO(p)).pages) for p in parts)


def scan_cost(pages):
    """Scans charged for a document: images count 1, PDFs 1 per started PAGES_PER_SCAN pages."""
    return max(1, -(-pages // PAGES_PER_SCAN))


def extract(data, mime, expense_cats, income_cats, api_key=None, model=None):
    """Return ScanResult as a dict for data already passed through prepare_upload().
    Raises RuntimeError with a Greek message on failure."""
    from google import genai
    from google.genai import types

    api_key = api_key or load_api_key()
    if not api_key:
        raise RuntimeError("Δεν βρέθηκε GEMINI_API_KEY (βάλε το στο αρχείο app\\.env)")
    client = genai.Client(api_key=api_key)
    prompt = PROMPT.format(
        today=date.today().isoformat(),
        expense_cats=", ".join(expense_cats) or "-",
        income_cats=", ".join(income_cats) or "-",
    )
    try:
        resp = client.models.generate_content(
            model=model or os.environ.get("GEMINI_MODEL", DEFAULT_MODEL),
            contents=([prompt, "Bank export (tab-separated):\n" + data.decode("utf-8")] if mime == "text/csv"
                      else [types.Part.from_bytes(data=data, mime_type=mime), prompt]),
            config=types.GenerateContentConfig(
                response_mime_type="application/json",
                response_schema=ScanResult,
                temperature=0,
            ),
        )
    except Exception as e:  # network, quota, bad key...
        raise RuntimeError(f"Σφάλμα Gemini: {str(e)[:200]}")
    parsed = resp.parsed
    if parsed is None:
        raise RuntimeError("Η απάντηση του Gemini δεν διαβάστηκε")
    return parsed.model_dump()


def _norm(s):
    """lowercase, no accents, no punctuation: 'Λογαριασμοί (ρεύμα/...)' -> 'λογαριασμοι ρευμα ...'"""
    s = unicodedata.normalize("NFD", s or "").lower()
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return " ".join("".join(ch if ch.isalnum() else " " for ch in s).split())


FALLBACK_NAMES = ("Άλλα", "Άλλο έσοδο")


def merchant_key(note):
    """Shop fingerprint for category memory: no case, accents, punctuation or number-only tokens
    ('ΣΚΛΑΒΕΝΙΤΗΣ #0452' == 'Σκλαβενίτης')."""
    return " ".join(w for w in _norm(note).split() if not any(ch.isdigit() for ch in w))


def same_merchant(a, b):
    """Same shop despite case/accents/branch numbers or extra words after the first two ('Shell Νίκαια 12')."""
    ka, kb = merchant_key(a), merchant_key(b)
    if not ka or not kb:
        return False
    if ka == kb:
        return True
    two_a, two_b = " ".join(ka.split()[:2]), " ".join(kb.split()[:2])
    return len(two_a) >= 6 and two_a == two_b


def recall(note, learned):
    """Category the user chose before for this shop: exact fingerprint, else same first two words."""
    key = merchant_key(note)
    if not key:
        return None
    if key in learned:
        return learned[key]
    two = " ".join(key.split()[:2])
    if len(two) >= 6:
        for k, cid in learned.items():                 # most recent first
            if " ".join(k.split()[:2]) == two:
                return cid
    return None


def match_category(name, kind, categories):
    """Exact name first, then tolerant match (model often shortens long names). Returns id or None."""
    same_kind = [c for c in categories if c["kind"] == kind]
    for c in same_kind:
        if c["name"] == name:
            return c["id"]
    n = _norm(name)
    if len(n) < 3:
        return None
    for c in same_kind:
        cn = _norm(c["name"])
        if cn == n or cn.startswith(n) or n.startswith(cn) or cn.split()[0] == n.split()[0]:
            return c["id"]
    return None


def clean_rows(result, categories, filename=""):
    """Validate model output against the real category list; returns rows ready for the review form."""
    fallback = {}
    for c in categories:
        if c["name"] in FALLBACK_NAMES:
            fallback[c["kind"]] = c["id"]
    rows = []
    for r in result.get("rows", []):
        if r.get("kind") == "transfer":
            continue                                   # own-account moves: neither income nor expense
        kind = r.get("kind") if r.get("kind") in ("income", "expense") else "expense"
        try:
            amount = abs(round(float(r.get("amount") or 0), 2))
        except (TypeError, ValueError):
            continue
        try:
            tdate = date.fromisoformat(str(r.get("tdate", ""))[:10]).isoformat()
        except ValueError:
            tdate = date.today().isoformat()  # visible on the review screen, user fixes it
        if amount <= 0:
            continue
        matched = match_category(r.get("category", ""), kind, categories)
        cat_id = matched or fallback.get(kind)
        rows.append({
            "tdate": tdate,
            "amount_cents": int(round(amount * 100)),
            "kind": kind,
            "note": (r.get("description") or "").strip()[:200],
            "category_id": cat_id,
            "guessed": bool(matched),
            "source": filename,
            "doc_type": result.get("document_type", ""),
        })
    return rows
