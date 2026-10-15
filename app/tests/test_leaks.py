"""User B tampers with every route using user A's ids. A's data must be unreadable and unchanged."""
import io

import pytest
from openpyxl import load_workbook

import repo
from conftest import cat_id

SECRET = "secret-A"


@pytest.fixture
def two(make_user, web, scope, ns):
    a, b = make_user("a"), make_user("b")
    with scope(a) as (cx, bid):
        cat = cat_id(cx, bid, "Καφέδες")
        tx = repo.add_transaction(cx, bid, "2026-09-10", cat, 777, f"{SECRET}-note")
        rec = repo.save_recurring(cx, bid, f"{SECRET}-rec", cat, 1234, 5)
        repo.save_category(cx, bid, f"{SECRET}-cat", "expense", 100, "#123456", "🔒")
        b_ = repo.get_budget(cx, bid)
        repo.update_budget(cx, bid, f"{SECRET}-title", b_["currency"], b_["accent"], 4242)
    return ns(a_id=a, b_id=b, a=web(a), b=web(b), cat=cat, tx=tx, rec=rec, bid=bid)


def a_state(scope, t):
    with scope(t.a_id) as (cx, bid):
        return (repo.get_transaction(cx, bid, t.tx),
                sorted((dict(c) for c in repo.list_categories(cx, bid, include_inactive=True)), key=lambda c: c["name"]),
                [dict(r) for r in repo.list_recurring(cx, bid)],
                dict(repo.get_budget(cx, bid)),
                len(repo.list_transactions(cx, bid)))


@pytest.mark.parametrize("url", ["/", "/?m=2026-09", "/tx", "/tx?m=2026-09", "/tx?q=secret",
                                 "/categories", "/recurring", "/settings", "/scan",
                                 "/export.csv", "/export.csv?m=2026-09"])
def test_b_pages_contain_nothing_of_a(two, url):
    r = two.b.get(url)
    assert r.status_code == 200
    assert SECRET.encode() not in r.data


def test_a_sees_own_secret(two):
    """Control: the marker really is on A's pages, so the checks above are meaningful."""
    assert SECRET.encode() in two.a.get("/tx?m=2026-09").data
    assert SECRET.encode() in two.a.get("/categories").data


def test_b_filter_by_a_category_shows_nothing(two):
    r = two.b.get(f"/tx?cat={two.cat}")
    assert r.status_code == 200 and SECRET.encode() not in r.data


def test_b_xlsx_export_contains_nothing_of_a(two):
    for url in ("/export.xlsx", "/export.xlsx?m=2026-09"):
        wb = load_workbook(io.BytesIO(two.b.get(url).data))
        cells = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]
        assert not any(SECRET in c for c in cells)


def test_b_cannot_read_a_transaction(two):
    assert two.b.get(f"/tx/{two.tx}/edit").status_code == 404


def test_b_tampering_changes_nothing(two, scope):
    before = a_state(scope, two)
    b = two.b
    assert b.post(f"/tx/{two.tx}/edit", data={"tdate": "2026-09-11", "category_id": two.cat,
                                               "amount": "1", "note": "hacked"}).status_code == 404
    assert b.post(f"/tx/{two.tx}/delete").status_code == 404
    b.post("/tx/add", data={"tdate": "2026-09-10", "category_id": two.cat, "amount": "5", "note": "hacked"})
    b.post("/categories", data={"id": two.cat, "name": "hacked", "kind": "expense", "active": "1"})
    b.post(f"/categories/{two.cat}/delete", data={})
    b.post("/recurring", data={"id": two.rec, "name": "hacked", "category_id": two.cat, "amount": "1", "day": "1"})
    assert b.post(f"/recurring/{two.rec}/delete").status_code == 404
    b.post("/recurring/post", data={"m": "2026-10", "ids": [str(two.rec)]})
    b.post("/scan/import", data={"count": 1, "r0_use": "1", "r0_date": "2026-09-10", "r0_amount": "9",
                                 "r0_cat": two.cat, "r0_note": "hacked"})
    b.post("/settings", data={"title": "hacked", "currency": "X", "accent": "#000000", "opening_balance": "1"})
    assert a_state(scope, two) == before


def test_b_own_budget_unaffected_by_failed_tampering(two, scope):
    two.b.post("/tx/add", data={"tdate": "2026-09-10", "category_id": two.cat, "amount": "5"})
    with scope(two.b_id) as (cx, bid):
        assert bid != two.bid
        assert repo.list_transactions(cx, bid) == []
