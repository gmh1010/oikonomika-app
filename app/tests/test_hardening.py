"""Browser-, upload-, export- and logging-level hardening (spec §6)."""
import csv
import io
import logging
import re
from datetime import date

import pytest
from PIL import Image

import repo
import scan
from conftest import APP_URL, LOCAL_AUTH, cat_id

PAGES = ["/", "/tx", "/categories", "/recurring", "/settings", "/scan"]


# ---------- headers & CSP ----------

def csp_parts(resp):
    return {p.strip().split(" ", 1)[0]: p.strip() for p in resp.headers["Content-Security-Policy"].split(";") if p.strip()}


def test_security_headers_on_app_pages(make_user, web):
    c = web(make_user("h"))
    for url in PAGES:
        r = c.get(url)
        csp = csp_parts(r)
        assert csp["script-src"] == "script-src 'self' https://challenges.cloudflare.com", url
        assert csp["default-src"] == "default-src 'self'"
        for d in ("frame-ancestors 'none'", "object-src 'none'", "base-uri 'none'", "form-action 'self'"):
            assert d in r.headers["Content-Security-Policy"], d
        assert r.headers["X-Content-Type-Options"] == "nosniff"
        assert r.headers["Referrer-Policy"] == "strict-origin-when-cross-origin"
        assert "camera=(self)" in r.headers["Permissions-Policy"] and "geolocation=()" in r.headers["Permissions-Policy"]
        assert r.headers["Cross-Origin-Opener-Policy"] == "same-origin"
        assert r.headers["X-Frame-Options"] == "DENY"
        assert "no-store" in r.headers["Cache-Control"], url
        assert r.headers["Strict-Transport-Security"] == "max-age=31536000; includeSubDomains"  # secure default


def test_hsts_only_when_secure(pool):
    from cryptography.fernet import Fernet

    from app import create_app

    def login_page(secure):
        app = create_app({"DATABASE_URL": APP_URL, "POOL": pool, "SECRET_KEY": "t", "TESTING": True, "IP_SALT": "s",
                          "SESSION_ENC_KEY": Fernet.generate_key().decode(), "COOKIE_SECURE": secure, **LOCAL_AUTH})
        return app.test_client().get("/auth/login", base_url="https://localhost" if secure else "http://localhost")
    assert login_page(True).headers["Strict-Transport-Security"] == "max-age=31536000; includeSubDomains"
    assert "Strict-Transport-Security" not in login_page(False).headers   # local plain-http dev only


def test_no_inline_scripts_anywhere(make_user, web):
    c = web(make_user("h"))
    for url in PAGES + ["/?m=2026-09"]:
        html = c.get(url).get_data(as_text=True)
        inline = re.findall(r"<script(?![^>]*\bsrc=)[^>]*>", html)
        assert inline == [], (url, inline)
        assert re.search(r"\son[a-z]+\s*=", html) is None, url      # no onclick= etc.
        assert "javascript:" not in html


def test_static_js_served(make_user, web):
    r = web(make_user("h")).get("/static/app.js")
    assert r.status_code == 200 and "data-kind-form" in r.get_data(as_text=True)


def test_templates_never_disable_escaping():
    import os
    root = os.path.join(os.path.dirname(os.path.dirname(__file__)), "templates")
    for dirpath, _, files in os.walk(root):
        for f in files:
            src = open(os.path.join(dirpath, f), encoding="utf-8").read()
            assert "|safe" not in src and "autoescape false" not in src, f


# ---------- uploads ----------

def _img(fmt, size=(40, 30), **save):
    b = io.BytesIO()
    Image.new("RGB", size, "white").save(b, fmt, **save)
    return b.getvalue()


def _pdf(pages):
    from pypdf import PdfWriter
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=200, height=200)
    b = io.BytesIO()
    w.write(b)
    return b.getvalue()


@pytest.mark.parametrize("fmt", ["PNG", "JPEG", "WEBP"])
def test_images_accepted_by_content_and_reencoded(fmt):
    data, out_mime = scan.prepare_upload(_img(fmt))
    assert out_mime == "image/jpeg" and data[:3] == b"\xff\xd8\xff"


def test_pdf_accepted_and_page_limit():
    data, mime = scan.prepare_upload(_pdf(scan.MAX_PDF_PAGES))
    assert mime == "application/pdf"
    with pytest.raises(ValueError, match="σελίδες"):
        scan.prepare_upload(_pdf(scan.MAX_PDF_PAGES + 1))


def test_long_pdf_split_into_small_chunks():
    from pypdf import PdfReader
    parts = scan.split_pdf(_pdf(32))
    sizes = [len(PdfReader(io.BytesIO(p)).pages) for p in parts]
    assert sizes == [5, 5, 5, 5, 5, 5, 2]
    assert scan.split_pdf(_pdf(3)) and len(scan.split_pdf(_pdf(3))) == 1
    assert scan.scan_cost(1) == 1 and scan.scan_cost(10) == 1 and scan.scan_cost(11) == 2 and scan.scan_cost(32) == 4


def _xlsx(n_rows, preamble=True):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    if preamble:
        ws.append(["Piraeus Bank"])
        ws.append(["Global Transaction List"])
        ws.append([])
    ws.append(["Category", "Transaction Narrative", "Posted date", "Amount", "Currency"])
    for i in range(n_rows):
        ws.append(["Services", f"SHOP {i}", "01/09/2026", -1.5 - i, "EUR"])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_spreadsheet_becomes_text_with_header_in_every_chunk():
    data, mime = scan.prepare_upload(_xlsx(450), "Global Transaction List.xlsx")
    assert mime == "text/csv"
    text = data.decode("utf-8")
    assert "Piraeus Bank" in text and "SHOP 449" in text and "\t" in text
    parts = scan.split_table(data)
    assert len(parts) == 3                                           # 450 rows -> 3 chunks of <= 200
    for p in parts:
        assert "Transaction Narrative" in p.decode("utf-8")          # header repeated so each chunk makes sense
    assert sum(p.decode("utf-8").count("SHOP ") for p in parts) == 450
    assert scan.table_rows(parts) == 450 and scan.scan_cost_rows(450) == 3 and scan.scan_cost_rows(12) == 1


def test_greek_csv_in_windows_encoding():
    raw = "Ημερομηνία;Περιγραφή;Ποσό\n01/09/2026;ΣΚΛΑΒΕΝΙΤΗΣ;-12,50\n".encode("cp1253")
    data, mime = scan.prepare_upload(raw, "kiniseis.csv")
    assert mime == "text/csv" and "ΣΚΛΑΒΕΝΙΤΗΣ" in data.decode("utf-8")


def test_bad_spreadsheets_rejected():
    import zipfile
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:            # zip bomb shaped like an .xlsx
        z.writestr("[Content_Types].xml", "<x/>")
        z.writestr("xl/workbook.xml", "<x/>")
        z.writestr("xl/worksheets/sheet1.xml", "0" * (80 * 1024 * 1024))
    with pytest.raises(ValueError, match="μεγάλο"):
        scan.prepare_upload(b.getvalue(), "bomb.xlsx")
    with pytest.raises(ValueError, match="xlsx"):
        scan.prepare_upload(bytes.fromhex("d0cf11e0a1b11ae1") + b"\x00" * 600, "old.xls")
    with pytest.raises(ValueError):
        scan.prepare_upload(_xlsx(scan.MAX_TABLE_ROWS + 1, preamble=False), "huge.xlsx")
    with pytest.raises(ValueError):
        scan.prepare_upload(bytes(range(256)) * 10, "binary.csv")      # not text


@pytest.mark.parametrize("data", [b"hello, I am text pretending to be a PNG", b"\x89PNG\r\n\x1a\nbroken",
                                  b"\x00\x00\x00\x18ftypheic" + b"\x00" * 40, b"%PDF-1.7 garbage", b""])
def test_bad_or_unsupported_files_rejected(data):
    with pytest.raises(ValueError):
        scan.prepare_upload(data)


def test_too_big_rejected():
    with pytest.raises(ValueError, match="10 MB"):
        scan.prepare_upload(b"\xff\xd8\xff" + b"\x00" * (10 * 1024 * 1024 + 1))


def test_exif_gps_is_stripped():
    img = Image.new("RGB", (60, 40), "white")
    exif = Image.Exif()
    exif[0x010F] = "PhoneMaker"                                    # Make
    exif[0x8825] = {1: "N", 2: (37.0, 58.0, 0.0), 3: "E", 4: (23.0, 43.0, 0.0)}   # GPS IFD
    b = io.BytesIO()
    img.save(b, "JPEG", exif=exif)
    assert Image.open(io.BytesIO(b.getvalue())).getexif()          # sanity: the source has EXIF
    out, _ = scan.prepare_upload(b.getvalue())
    cleaned = Image.open(io.BytesIO(out))
    assert not cleaned.getexif() and b"PhoneMaker" not in out and "exif" not in cleaned.info


def test_big_photo_downscaled():
    out, _ = scan.prepare_upload(_img("PNG", size=(4000, 3000)))
    assert max(Image.open(io.BytesIO(out)).size) == scan.MAX_SIDE


def test_fake_file_does_not_use_quota(make_user, web, scope):
    u = make_user("q")
    c = web(u)
    calls = []
    c.application.config["EXTRACTOR"] = lambda *a: calls.append(1) or {"document_type": "other", "rows": []}
    r = c.post("/scan", data={"files": (io.BytesIO(b"not an image at all"), "evil.png", "image/png")},
               content_type="multipart/form-data")
    assert "δεν είναι" in r.get_data(as_text=True) and calls == []
    with scope(u) as (cx, bid):
        assert repo.scans_used(cx, u, date.today().strftime("%Y-%m")) == 0


# ---------- exports ----------

# Spreadsheet formula starters that must never reach Excel/LibreOffice as live formulas.
EVIL = ['=HYPERLINK("http://example.invalid","x")', "+1+1", "-2+3", "@SUM(A1)", "\t=1", "\r=1"]
EVIL_CAT = "=SUM(1,2)"


def test_formula_injection_neutralised_in_exports(make_user, web, scope):
    u = make_user("x")
    with scope(u) as (cx, bid):
        coffee = cat_id(cx, bid, "Καφέδες")
        for i, note in enumerate(EVIL):
            repo.add_transaction(cx, bid, f"2026-09-{10 + i:02d}", coffee, 100 + i, note)
        repo.add_transaction(cx, bid, "2026-09-20", coffee, 999, "normal note")
        repo.save_category(cx, bid, EVIL_CAT, "expense", 0, "#123456", "")
        repo.add_transaction(cx, bid, "2026-09-21", cat_id(cx, bid, EVIL_CAT), 5, "cat test")
    c = web(u)
    rows = list(csv.reader(io.StringIO(c.get("/export.csv?m=2026-09").data.decode("utf-8-sig")), delimiter=";"))
    cells = [cell for row in rows for cell in row]
    for cell in cells:
        assert not cell.startswith(("=", "+", "-", "@", "\t", "\r")), cell
    assert "normal note" in cells and "'" + EVIL_CAT in cells and "'+1+1" in cells
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(c.get("/export.xlsx?m=2026-09").data))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", cell.value
                if isinstance(cell.value, str):
                    assert not cell.value.startswith(("=", "+", "-", "@", "\t", "\r")), cell.value
    notes = [r[4].value for r in wb["Κινήσεις"].iter_rows(min_row=2)]
    assert "normal note" in notes and any(n.startswith("'=HYPERLINK") for n in notes)
    amounts = [r[3].value for r in wb["Κινήσεις"].iter_rows(min_row=2)]
    assert all(isinstance(a, (int, float)) for a in amounts)       # numbers stay numbers


# ---------- errors & logging ----------

def test_generic_error_pages(make_user, web):
    c = web(make_user("e"))
    r = c.get("/definitely-not-here")
    assert r.status_code == 404 and "Δεν βρέθηκε" in r.get_data(as_text=True)
    assert c.put("/").status_code == 405


def test_500_page_hides_internals(make_user, pool):
    from app import create_app
    u = make_user("e")
    app = create_app({"DATABASE_URL": APP_URL, "POOL": pool, "SECRET_KEY": "t", "TESTING": True,
                      "USER_RESOLVER": lambda: u})
    app.config["PROPAGATE_EXCEPTIONS"] = False

    @app.route("/boom")
    def boom():
        raise RuntimeError("SELECT secret_column FROM oiko.transactions -- internal detail")
    r = app.test_client().get("/boom")
    body = r.get_data(as_text=True)
    assert r.status_code == 500 and "Κάτι πήγε στραβά" in body
    assert "secret_column" not in body and "Traceback" not in body


def test_413_page(make_user, web):
    c = web(make_user("e"))
    c.application.config["MAX_CONTENT_LENGTH"] = 1000
    r = c.post("/scan", data={"files": (io.BytesIO(b"x" * 5000), "big.png", "image/png")},
               content_type="multipart/form-data")
    assert r.status_code == 413 and "μεγάλο" in r.get_data(as_text=True)


def test_request_log_has_no_query_or_secrets(make_user, web, caplog):
    c = web(make_user("l"))
    caplog.set_level(logging.INFO, logger="oiko.request")
    c.get("/tx?q=my-secret-search&m=2026-09")
    c.get("/auth/confirm?token_hash=deadbeefdeadbeefdeadbeefdeadbeef&type=email")
    lines = [r.getMessage() for r in caplog.records if r.name == "oiko.request"]
    assert any('"path": "/tx"' in line and '"status": 200' in line for line in lines), lines
    joined = "\n".join(lines)
    assert "my-secret-search" not in joined and "deadbeef" not in joined and "?" not in joined


def test_proxyfix_only_when_trusted(pool, make_user):
    from app import create_app
    u = make_user("p")
    seen = {}

    def make(trust):
        app = create_app({"DATABASE_URL": APP_URL, "POOL": pool, "SECRET_KEY": "t", "TESTING": True,
                          "USER_RESOLVER": lambda: u, "TRUST_PROXY": trust})

        @app.route("/_ip")
        def _ip():
            from flask import request
            seen[trust] = request.remote_addr
            return "ok"
        return app.test_client()
    make(False).get("/_ip", headers={"X-Forwarded-For": "6.6.6.6"})
    make(True).get("/_ip", headers={"X-Forwarded-For": "6.6.6.6"})
    assert seen[False] != "6.6.6.6" and seen[True] == "6.6.6.6"
