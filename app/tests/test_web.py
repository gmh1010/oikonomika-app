import io

import repo
from conftest import cat_id


def test_pages_render(make_user, web):
    c = web(make_user("w"))
    for url in ["/", "/?m=2026-09", "/?m=bad", "/tx", "/tx?m=2026-09&kind=expense&q=x",
                "/categories", "/recurring", "/settings", "/scan"]:
        assert c.get(url).status_code == 200, url
    assert "Σεπτέμβριος 2026" in c.get("/?m=2026-09").get_data(as_text=True)


def test_no_user_is_401(pool):
    from app import create_app
    from conftest import APP_URL
    app = create_app({"DATABASE_URL": APP_URL, "POOL": pool, "SECRET_KEY": "t", "TESTING": True,
                      "USER_RESOLVER": lambda: None})
    assert app.test_client().get("/").status_code == 401


def test_add_edit_delete_flow(make_user, web, scope):
    u = make_user("w")
    c = web(u)
    with scope(u) as (cx, bid):
        coffee = cat_id(cx, bid, "Καφέδες")
    r = c.post("/tx/add", data={"tdate": "2026-09-12", "category_id": coffee, "amount": "3,20",
                                "note": "freddo", "next": "/?m=2026-09"})
    assert r.status_code == 302 and r.headers["Location"].endswith("/?m=2026-09")
    page = c.get("/?m=2026-09").get_data(as_text=True)
    assert "freddo" in page and "3,20" in page
    with scope(u) as (cx, bid):
        tx = repo.list_transactions(cx, bid)[0]["id"]
    assert c.get(f"/tx/{tx}/edit").status_code == 200
    c.post(f"/tx/{tx}/edit", data={"tdate": "2026-09-12", "category_id": coffee, "amount": "4", "note": "freddo x2"})
    assert "freddo x2" in c.get("/tx?m=2026-09").get_data(as_text=True)
    c.post(f"/tx/{tx}/delete")
    assert "freddo" not in c.get("/tx?m=2026-09").get_data(as_text=True)
    assert c.get("/tx/not-a-uuid/edit").status_code == 404


def test_bad_input_flashes_and_does_not_crash(make_user, web, scope):
    u = make_user("w")
    c = web(u)
    with scope(u) as (cx, bid):
        other = cat_id(cx, bid, "Άλλα")
    r = c.post("/tx/add", data={"tdate": "2026-09-12", "category_id": other, "amount": "abc"}, follow_redirects=True)
    assert "Μη έγκυρο ποσό" in r.get_data(as_text=True)
    r = c.post("/tx/add", data={"tdate": "2026-09-12", "category_id": other, "amount": "inf"}, follow_redirects=True)
    assert "Μη έγκυρο ποσό" in r.get_data(as_text=True)
    r = c.post("/tx/add", data={"tdate": "2026-09-12", "category_id": "' or 1=1 --", "amount": "1"},
               follow_redirects=True)
    assert r.status_code == 200 and "Μη έγκυρο" in r.get_data(as_text=True)
    with scope(u) as (cx, bid):
        assert repo.list_transactions(cx, bid) == []


def test_open_redirect_blocked(make_user, web, scope):
    u = make_user("w")
    with scope(u) as (cx, bid):
        other = cat_id(cx, bid, "Άλλα")
    for evil in ("//evil.example", "/\\evil.example"):
        r = web(u).post("/tx/add", data={"tdate": "2026-09-12", "category_id": other, "amount": "1", "next": evil})
        assert "evil" not in r.headers["Location"]


def test_customisation(make_user, web):
    c = web(make_user("w"))
    c.post("/categories", data={"name": "Γυμναστήριο", "kind": "expense", "budget": "35",
                                "color": "#ff0000", "icon": "🏋️"})
    page = c.get("/categories").get_data(as_text=True)
    assert "Γυμναστήριο" in page and 'value="🏋️"' in page
    r = c.post("/categories", data={"name": "Γυμναστήριο", "kind": "expense"}, follow_redirects=True)
    assert "Υπάρχει ήδη" in r.get_data(as_text=True)
    c.post("/settings", data={"title": "Τα λεφτά μου", "currency": "EUR", "accent": "#123abc",
                              "opening_balance": "-120,50"})
    page = c.get("/").get_data(as_text=True)
    assert "Τα λεφτά μου" in page and "#123abc" in page and "-120,50 EUR" in page
    assert "🛒 Σούπερ μάρκετ" in page


def test_post_recurring(make_user, web, scope):
    u = make_user("w")
    with scope(u) as (cx, bid):
        repo.save_recurring(cx, bid, "Ενοίκιο", cat_id(cx, bid, "Ενοίκιο"), 30000, 1)
    r = web(u).post("/recurring/post", data={"m": "2026-09"}, follow_redirects=True)
    assert "Περάστηκαν 1 πάγια" in r.get_data(as_text=True)


def test_csrf_enforced(make_user, web):
    r = web(make_user("w"), csrf=True).post("/tx/add", data={"amount": "1"})
    assert r.status_code == 400


def test_test_login_hook_refused_outside_testing(pool):
    import pytest
    from app import create_app
    from conftest import APP_URL
    with pytest.raises(RuntimeError, match="test hook"):
        create_app({"DATABASE_URL": APP_URL, "POOL": pool, "SECRET_KEY": "t", "USER_RESOLVER": lambda: "x"})


def test_real_config_requires_auth_keys(pool):
    import pytest
    from app import create_app
    from conftest import APP_URL
    with pytest.raises(RuntimeError, match="missing config"):
        create_app({"DATABASE_URL": APP_URL, "POOL": pool, "SECRET_KEY": "t"})


def test_exports(make_user, web, scope):
    u = make_user("w")
    with scope(u) as (cx, bid):
        repo.add_transaction(cx, bid, "2026-09-10", cat_id(cx, bid, "Καφέδες"), 250, "freddo")
    c = web(u)
    r = c.get("/export.xlsx?m=2026-09")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.properties.creator == "Oikonomika"
    assert wb.sheetnames == ["Κινήσεις", "Σύνοψη"] and wb["Κινήσεις"].max_row == 2
    assert "freddo" in c.get("/export.csv?m=2026-09").data.decode("utf-8-sig")


def test_make_recurring_button(make_user, web, scope):
    a, b = make_user("ra"), make_user("rb")
    with scope(a) as (cx, bid):
        tx = repo.add_transaction(cx, bid, "2026-09-03", cat_id(cx, bid, "Ενοίκιο"), 30000, "ΖΗΚΟΥΛΗΣ")
    ca = web(a)
    assert "Κάν’ το πάγιο" in ca.get(f"/tx/{tx}/edit").get_data(as_text=True)
    assert web(b).post(f"/tx/{tx}/recurring").status_code == 404        # someone else's transaction
    r = ca.post(f"/tx/{tx}/recurring", follow_redirects=True)
    page = r.get_data(as_text=True)
    assert "Έγινε πάγιο" in page and "κάθε μήνα στις 3" in page
    assert "Είναι πάγιο" in ca.get(f"/tx/{tx}/edit").get_data(as_text=True)
    with scope(b) as (cx, bid):
        assert repo.list_recurring(cx, bid) == []
