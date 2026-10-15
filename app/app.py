"""Oikonomika — multi-user expense / income tracker (Flask + Postgres with Row-Level Security)."""
import csv
import io
import re
import secrets
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from types import SimpleNamespace

from flask import (Flask, abort, flash, g, redirect, render_template, request,
                   send_file, session, url_for)

import auth
import db
import hardening
import repo
import scan
import security
import ui
from authstore import AuthStore
from gotrue import GoTrue
from money import fmt, to_cents

MONTH_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
COLOR_RE = re.compile(r"^#[0-9a-fA-F]{6}$")
EMOJIS = ("💼 💻 💰 💶 🎁 🏠 🏦 💳 💡 🔥 💧 📱 🌐 🛒 🍔 🍕 ☕ 🍺 🚬 ⛽ 🏍️ 🚗 🚌 🅿️ 📺 🎮 🎬 🎵 "
          "👕 👟 💊 🏥 🦷 💈 🧴 🏋️ ⚽ ✈️ ⛴️ 🏖️ 🐶 👶 🎓 📚 🧾 🏛️ 🔧 🛠️ 📦 ❤️ 🎉").split()
MONTHS_GR = ["Ιανουάριος", "Φεβρουάριος", "Μάρτιος", "Απρίλιος", "Μάιος", "Ιούνιος",
             "Ιούλιος", "Αύγουστος", "Σεπτέμβριος", "Οκτώβριος", "Νοέμβριος", "Δεκέμβριος"]
MONTHS_SHORT = ["Ιαν", "Φεβ", "Μάρ", "Απρ", "Μάι", "Ιούν", "Ιούλ", "Αύγ", "Σεπ", "Οκτ", "Νοέ", "Δεκ"]
MAX_IMPORT_ROWS = 1000               # rows on one scan review / import (a 40-page statement fits)
DEFAULT_S = {"title": "Οικονομικά", "currency": "€", "accent": "#2563eb", "opening_balance_cents": 0, "theme": "auto"}
FORMULA_START = ("=", "+", "-", "@", chr(9), chr(13))   # = + - @ TAB CR


def safe_cell(value):
    """Neutralise spreadsheet formula injection: text that Excel/LibreOffice would run becomes plain text."""
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return value


AUTH_KEYS = ("SUPABASE_URL", "SUPABASE_PUBLISHABLE_KEY", "SUPABASE_SECRET_KEY", "SESSION_ENC_KEY", "IP_SALT")


def create_app(config):
    """Required: DATABASE_URL, SECRET_KEY and the AUTH_KEYS. Optional: POOL, TESTING, EXTRACTOR, SCAN_QUOTA,
    COOKIE_SECURE (default True), INVITE_CODE, TURNSTILE_SITE_KEY/TURNSTILE_SECRET, GOOGLE_ENABLED, PWNED_CHECK,
    CSRF (default: on unless TESTING), TRUST_PROXY (behind one reverse proxy, e.g. Caddy).
    Tests only: USER_RESOLVER (callable -> user uuid) replaces login entirely and is refused outside TESTING."""
    app = Flask(__name__)
    app.config.update(MAX_CONTENT_LENGTH=40 * 1024 * 1024, MAX_FORM_PARTS=MAX_IMPORT_ROWS * 6 + 20,
                      MAX_FORM_MEMORY_SIZE=4 * 1024 * 1024, EXTRACTOR=scan.extract, TESTING=False, SCAN_QUOTA=30,
                      COOKIE_SECURE=True, GOOGLE_ENABLED=False, PWNED_CHECK=security.pwned, TRUST_PROXY=False)
    app.config.update(config)
    resolver = app.config.get("USER_RESOLVER")
    if resolver and not app.config["TESTING"]:
        raise RuntimeError("USER_RESOLVER is a test hook and cannot be used outside TESTING")
    for key in ("DATABASE_URL", "SECRET_KEY") + (() if resolver else AUTH_KEYS):
        if not app.config.get(key):
            raise RuntimeError(f"missing config: {key}")
    secure = app.config["COOKIE_SECURE"]
    app.secret_key = app.config["SECRET_KEY"]
    app.config.update(SESSION_COOKIE_NAME="__Host-oiko_f" if secure else "oiko_f", SESSION_COOKIE_SECURE=secure,
                      SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Lax")
    pool = app.config.get("POOL") or db.make_pool(app.config["DATABASE_URL"])
    app.extensions["oiko"] = SimpleNamespace(
        pool=pool,
        gotrue=None if resolver else GoTrue(app.config["SUPABASE_URL"], app.config["SUPABASE_PUBLISHABLE_KEY"],
                                            app.config["SUPABASE_SECRET_KEY"]),
        store=None if resolver else AuthStore(pool, app.config["SESSION_ENC_KEY"]),
    )
    app.register_blueprint(auth.bp)
    hardening.init_app(app)

    # ---------- request scope: one transaction per request, bound to the user ----------

    def csrf_token():
        if "csrf" not in session:
            session["csrf"] = secrets.token_hex(16)
        return session["csrf"]

    @app.before_request
    def check_csrf():
        if request.method == "POST" and app.config.get("CSRF", not app.config["TESTING"]):
            expected = session.get("csrf")
            if not expected or not secrets.compare_digest(request.form.get("csrf", ""), expected):
                abort(400, "Λήξη φόρμας — ανανέωσε τη σελίδα")

    @app.before_request
    def authenticate():
        ep = request.endpoint
        if ep in (None, "static"):
            return None
        if resolver:
            uid = resolver()
            g.user = {"user_id": uid, "email": "", "mfa_pending": False} if uid else None
        else:
            g.user = auth.load_user()
        if ep in auth.PUBLIC or ep == "healthz":
            return None
        if not g.user:
            if request.method in ("GET", "HEAD") and not resolver:
                return redirect(url_for("auth.login", next=request.full_path.rstrip("?")))
            abort(401)
        if g.user["mfa_pending"] and ep not in auth.MFA_PENDING_OK:
            return redirect(url_for("auth.mfa"))
        return None

    @app.before_request
    def open_scope():
        if request.endpoint in (None, "static", "healthz") or request.endpoint.startswith("auth."):
            return None
        g.uid = repo.uid(g.user["user_id"])
        g._scope = db.scoped(pool, g.uid)
        g.cx = g._scope.__enter__()
        g.bid = repo.provision_budget(g.cx)
        return None

    @app.teardown_request
    def close_scope(exc):
        scope = g.pop("_scope", None)
        if scope is not None:
            if exc is None:
                scope.__exit__(None, None, None)
            else:
                scope.__exit__(type(exc), exc, exc.__traceback__)

    @app.route("/healthz")
    def healthz():
        """Liveness + database reachability for the container health check and uptime monitor."""
        try:
            with pool.connection() as c:
                c.execute("select 1")
        except Exception:
            return {"ok": False}, 503
        return {"ok": True}

    # ---------- helpers ----------

    def current_month():
        m = request.values.get("m") or ""
        return m if MONTH_RE.match(m) else date.today().strftime("%Y-%m")

    def shift_month(month, delta):
        y, m = map(int, month.split("-"))
        m += delta
        while m > 12:
            y, m = y + 1, m - 12
        while m < 1:
            y, m = y - 1, m + 12
        return f"{y:04d}-{m:02d}"

    def month_label(month):
        y, m = map(int, month.split("-"))
        return f"{MONTHS_GR[m - 1]} {y}"

    def optional_uuid(value):
        try:
            return repo.uid(value) if value else None
        except ValueError:
            return None

    @app.context_processor
    def inject():
        S = dict(DEFAULT_S)
        sheet_cats = []
        b = None
        if "cx" in g:
            b = repo.get_budget(g.cx, g.bid)
            sheet_cats = repo.list_categories(g.cx, g.bid)
        elif g.get("user") and not g.user["mfa_pending"]:
            # auth.* pages open no request scope; still show them in the user's own theme/accent
            with db.scoped(pool, repo.uid(g.user["user_id"])) as cx:
                b = repo.get_budget(cx, repo.provision_budget(cx))
        if b:
            S = {"title": b["name"], "currency": b["currency"], "accent": b["accent"],
                 "opening_balance_cents": b["opening_balance_cents"], "theme": b["theme"]}
        cur = S["currency"]
        user = g.get("user")
        email = (user or {}).get("email") or ""
        return {"S": S, "money": lambda c: fmt(c, cur), "plain": lambda c: fmt(c, "").strip(),
                "csrf_token": csrf_token, "month_label": month_label, "shift_month": shift_month,
                "EMOJIS": EMOJIS, "user": user, "TURNSTILE_SITE_KEY": app.config.get("TURNSTILE_SITE_KEY"),
                "GOOGLE_ENABLED": app.config["GOOGLE_ENABLED"], "icon": ui.icon,
                "accent_rgb": ui.hex_to_rgb(S["accent"]),
                "theme_class": {"light": "light", "dark": "dark"}.get(S.get("theme"), ""),
                "sheet_cats": sheet_cats, "today_iso": date.today().isoformat(),
                "user_email": email, "user_initials": (email[:2] or "··").upper(),
                "is_on": lambda *eps: request.endpoint in eps,
                "month_short": lambda m: MONTHS_SHORT[int(m[5:7]) - 1]}

    def form_amount(field="amount"):
        try:
            return to_cents(request.form.get(field, ""))
        except ValueError:
            raise ValueError("Μη έγκυρο ποσό")

    def back(default="dashboard", **kw):
        nxt = request.form.get("next") or ""
        if nxt.startswith("/") and not nxt.startswith("//") and "\\" not in nxt:
            return redirect(nxt)
        return redirect(url_for(default, **kw))

    # ---------- dashboard ----------

    @app.route("/")
    def dashboard():
        month = current_month()
        cx, bid = g.cx, g.bid
        start, end = repo.month_bounds(month)
        trend = repo.last_months(cx, bid, month, 6)
        peak = max([max(t["income"], t["expense"]) for t in trend] + [1])
        today = date.today()
        is_current = today.strftime("%Y-%m") == month
        days_in = (date.fromisoformat(end) - date.fromisoformat(start)).days
        return render_template(
            "dashboard.html", month=month, s=repo.month_summary(cx, bid, month),
            day_of_month=today.day if is_current else None, days_in_month=days_in,
            balance=repo.balance_until(cx, bid, end), pending=repo.pending_recurring(cx, bid, month),
            recent=repo.list_transactions(cx, bid, month=month, limit=8), trend=trend, peak=peak,
            cats_income=repo.list_categories(cx, bid, "income"),
            cats_expense=repo.list_categories(cx, bid, "expense"),
            today=date.today().isoformat() if date.today().strftime("%Y-%m") == month else start)

    # ---------- transactions ----------

    @app.route("/tx")
    def transactions():
        month = request.args.get("m", "")
        month = month if MONTH_RE.match(month) else ""
        kind = request.args.get("kind") if request.args.get("kind") in ("income", "expense") else None
        cat = optional_uuid(request.args.get("cat"))
        q = (request.args.get("q") or "").strip()[:60] or None
        rows = repo.list_transactions(g.cx, g.bid, month=month or None, kind=kind, category_id=cat, q=q)
        inc = sum(r["amount_cents"] for r in rows if r["kind"] == "income")
        exp = sum(r["amount_cents"] for r in rows if r["kind"] == "expense")
        return render_template("transactions.html", rows=rows, month=month, kind=kind, cat=cat, q=q or "",
                               inc=inc, exp=exp, cats=repo.list_categories(g.cx, g.bid, include_inactive=True))

    @app.route("/tx/add", methods=["POST"])
    def tx_add():
        try:
            repo.add_transaction(g.cx, g.bid, request.form.get("tdate", ""), request.form.get("category_id", ""),
                                 form_amount(), request.form.get("note", ""))
            flash("Καταχωρήθηκε ✓", "ok")
        except ValueError as e:
            flash(str(e) or "Μη έγκυρα στοιχεία", "err")
        return back()

    @app.route("/tx/<uuid:tx_id>/edit", methods=["GET", "POST"])
    def tx_edit(tx_id):
        tx = repo.get_transaction(g.cx, g.bid, tx_id)
        if not tx:
            abort(404)
        if request.method == "POST":
            try:
                if not repo.update_transaction(g.cx, g.bid, tx_id, request.form.get("tdate", ""),
                                               request.form.get("category_id", ""), form_amount(),
                                               request.form.get("note", "")):
                    abort(404)
                flash("Αποθηκεύτηκε ✓", "ok")
                return back("transactions")
            except ValueError as e:
                flash(str(e) or "Μη έγκυρα στοιχεία", "err")
        return render_template("tx_edit.html", tx=tx,
                               cats=repo.list_categories(g.cx, g.bid, include_inactive=True),
                               next=request.values.get("next", ""))

    @app.route("/tx/<uuid:tx_id>/recurring", methods=["POST"])
    def tx_make_recurring(tx_id):
        if not repo.get_transaction(g.cx, g.bid, tx_id):
            abort(404)
        try:
            rid = repo.make_recurring_from_tx(g.cx, g.bid, tx_id)
        except ValueError as e:
            flash(str(e), "err")
            return redirect(url_for("tx_edit", tx_id=tx_id))
        rec = next(r for r in repo.list_recurring(g.cx, g.bid) if r["id"] == rid)
        flash(f"Έγινε πάγιο: «{rec['name']}» κάθε μήνα στις {rec['day']} ✓", "ok")
        return redirect(url_for("recurring"))

    @app.route("/tx/<uuid:tx_id>/delete", methods=["POST"])
    def tx_delete(tx_id):
        if not repo.delete_transaction(g.cx, g.bid, tx_id):
            abort(404)
        flash("Διαγράφηκε", "ok")
        return back("transactions")

    # ---------- categories ----------

    @app.route("/categories", methods=["GET", "POST"])
    def categories():
        if request.method == "POST":
            try:
                color = request.form.get("color", "#6b7280")
                if not COLOR_RE.match(color):
                    color = "#6b7280"
                budget = request.form.get("budget", "").strip()
                try:
                    budget_cents = to_cents(budget) if budget else 0
                except ValueError:
                    raise ValueError("Μη έγκυρο ποσό")
                repo.save_category(g.cx, g.bid, request.form.get("name", "")[:60], request.form.get("kind", ""),
                                   budget_cents, color, request.form.get("icon", ""),
                                   request.form.get("active", "1") == "1", request.form.get("id") or None)
                flash("Η κατηγορία αποθηκεύτηκε ✓", "ok")
            except ValueError as e:
                flash(str(e), "err")
            return redirect(url_for("categories"))
        return render_template("categories.html", cats=repo.list_categories(g.cx, g.bid, include_inactive=True))

    @app.route("/categories/<uuid:cat_id>/delete", methods=["POST"])
    def category_delete(cat_id):
        try:
            repo.delete_category(g.cx, g.bid, cat_id, request.form.get("move_to") or None)
            flash("Η κατηγορία διαγράφηκε", "ok")
        except ValueError as e:
            flash(str(e), "err")
        return redirect(url_for("categories"))

    # ---------- recurring ----------

    @app.route("/recurring", methods=["GET", "POST"])
    def recurring():
        if request.method == "POST":
            try:
                repo.save_recurring(g.cx, g.bid, request.form.get("name", "")[:60],
                                    request.form.get("category_id", ""), form_amount(),
                                    request.form.get("day", 1, type=int),
                                    request.form.get("active", "1") == "1", request.form.get("id") or None)
                flash("Αποθηκεύτηκε ✓", "ok")
            except ValueError as e:
                flash(str(e), "err")
            return redirect(url_for("recurring"))
        month = date.today().strftime("%Y-%m")
        return render_template("recurring.html", items=repo.list_recurring(g.cx, g.bid),
                               cats=repo.list_categories(g.cx, g.bid), month=month,
                               posted=repo.recurring_posted(g.cx, g.bid, month))

    @app.route("/recurring/<uuid:rec_id>/delete", methods=["POST"])
    def recurring_delete(rec_id):
        if not repo.delete_recurring(g.cx, g.bid, rec_id):
            abort(404)
        flash("Διαγράφηκε", "ok")
        return redirect(url_for("recurring"))

    @app.route("/recurring/post", methods=["POST"])
    def recurring_post():
        month = current_month()
        ids = request.form.getlist("ids") or None
        try:
            n = repo.post_recurring(g.cx, g.bid, month, ids)
            flash(f"Περάστηκαν {n} πάγια για {month_label(month)} ✓", "ok")
        except ValueError as e:
            flash(str(e), "err")
        return redirect(url_for("dashboard", m=month))

    # ---------- scan ----------

    @app.route("/scan", methods=["GET", "POST"])
    def scan_page():
        if request.method in ("GET", "HEAD"):
            return render_template("scan.html", has_key=bool(scan.load_api_key()) or app.config["TESTING"],
                                   used=repo.scans_used(g.cx, g.uid, date.today().strftime("%Y-%m")),
                                   quota=app.config["SCAN_QUOTA"])
        files = [f for f in request.files.getlist("files") if f and f.filename]
        if not files:
            flash("Διάλεξε τουλάχιστον μία φωτογραφία ή PDF", "err")
            return redirect(url_for("scan_page"))
        month = date.today().strftime("%Y-%m")
        left = max(app.config["SCAN_QUOTA"] - repo.scans_used(g.cx, g.uid, month), 0)
        if left == 0:
            flash(f"Έφτασες το μηνιαίο όριο των {app.config['SCAN_QUOTA']} σαρώσεων", "err")
            return redirect(url_for("scan_page"))
        cats = [dict(c) for c in repo.list_categories(g.cx, g.bid)]
        exp_names = [c["name"] for c in cats if c["kind"] == "expense"]
        inc_names = [c["name"] for c in cats if c["kind"] == "income"]
        rows, errors_, used, jobs, transfers = [], [], 0, [], 0
        for f in files[:10]:
            name = (f.filename or "αρχείο")[:80]
            try:
                data, mime = scan.prepare_upload(f.read(scan.MAX_BYTES + 1), name)
                parts = (scan.split_pdf(data) if mime == "application/pdf" else
                         scan.split_table(data) if mime == "text/csv" else [data])
            except ValueError as e:
                errors_.append(f"{name}: {e}")
                continue
            except Exception:
                errors_.append(f"{name}: Το αρχείο δεν είναι έγκυρο PDF")
                continue
            cost = (scan.scan_cost(scan.pdf_pages(parts)) if mime == "application/pdf" else
                    scan.scan_cost_rows(scan.table_rows(parts)) if mime == "text/csv" else 1)
            if used + cost > left:
                errors_.append(f"{name}: έφτασες το μηνιαίο όριο σαρώσεων" if cost == 1 else
                               f"{name}: χρειάζεται {cost} σαρώσεις, σου μένουν {left - used} "
                               f"(όριο {app.config['SCAN_QUOTA']} τον μήνα)")
                continue
            used += cost
            jobs.append((name, mime, parts))

        extractor = app.config["EXTRACTOR"]

        def read_part(task):
            name, mime, part = task
            try:
                return extractor(part, mime, exp_names, inc_names), None
            except RuntimeError as e:
                return None, str(e)

        tasks = [(name, mime, part) for name, mime, parts in jobs for part in parts]
        with ThreadPoolExecutor(max_workers=8) as pool_:            # long statements: all chunks at once
            outcomes = iter(list(pool_.map(read_part, tasks)))
        for name, mime, parts in jobs:
            found, failed = [], []
            for i in range(len(parts)):
                result, err = next(outcomes)
                label = name if len(parts) == 1 else f"{name} (μέρος {i + 1}/{len(parts)})"
                if err:
                    failed.append(f"{label}: {err}")
                else:
                    transfers += sum(1 for r in result.get("rows", []) if r.get("kind") == "transfer")
                    found.extend(scan.clean_rows(result, cats, name))
            errors_.extend(failed)
            if not found and not failed and not transfers:
                errors_.append(f"{name}: δεν βρέθηκαν κινήσεις")
            rows.extend(found)
        if used:
            repo.add_scans(g.cx, g.uid, month, used)
        claimed, seen = set(), []
        for r in rows:                     # duplicate = same name + same date + same amount (+ direction)
            dup = repo.find_duplicate(g.cx, g.bid, r["tdate"], r["amount_cents"], r["kind"], r["note"], claimed)
            r["dup"] = dup
            if dup:
                claimed.add(dup["id"])
            r["dup_in_scan"] = not dup and any(
                (s["tdate"], s["amount_cents"], s["kind"]) == (r["tdate"], r["amount_cents"], r["kind"])
                and scan.same_merchant(s["note"], r["note"]) for s in seen)
            seen.append(r)
        rows.sort(key=lambda r: r["tdate"], reverse=True)
        rows = rows[:MAX_IMPORT_ROWS]
        learned = repo.learned_categories(g.cx, g.bid) if rows else {}
        kind_of = {c["id"]: c["kind"] for c in cats}
        for r in rows:                                   # the user's past choice beats the model's guess
            cid = scan.recall(r["note"], learned)
            if cid and kind_of.get(cid) == r["kind"]:
                r.update(category_id=cid, guessed=True, remembered=True)
        dups = sum(1 for r in rows if r["dup"] or r["dup_in_scan"])
        return render_template("scan_review.html", rows=rows, errors=errors_, transfers=transfers, dups=dups,
                               cats=repo.list_categories(g.cx, g.bid))

    @app.route("/scan/import", methods=["POST"])
    def scan_import():
        n, bad = 0, 0
        count = min(request.form.get("count", 0, type=int) or 0, MAX_IMPORT_ROWS)
        for i in range(count):
            if request.form.get(f"r{i}_use") != "1":
                continue
            try:
                repo.add_transaction(g.cx, g.bid, request.form.get(f"r{i}_date", ""),
                                     request.form.get(f"r{i}_cat", ""), to_cents(request.form.get(f"r{i}_amount", "")),
                                     request.form.get(f"r{i}_note", ""))
                n += 1
            except ValueError:
                bad += 1
        flash(f"Καταχωρήθηκαν {n} κινήσεις ✓" + (f" — {bad} με λάθη παραλείφθηκαν" if bad else ""),
              "ok" if not bad else "err")
        return redirect(url_for("transactions"))

    @app.route("/more")
    def more():
        return render_template("more.html")

    # ---------- settings ----------

    @app.route("/settings", methods=["GET", "POST"])
    def settings():
        if request.method == "POST":
            try:
                title = request.form.get("title", "").strip()[:40] or "Οικονομικά"
                currency = request.form.get("currency", "").strip()[:4] or "€"
                accent = request.form.get("accent", "#0f766e")
                if not COLOR_RE.match(accent):
                    accent = "#0f766e"
                ob = request.form.get("opening_balance", "").strip()
                neg = ob.startswith("-")
                opening = to_cents(ob.lstrip("-")) if ob.lstrip("-") else 0
                theme = request.form.get("theme", "auto")
                theme = theme if theme in ("auto", "light", "dark") else "auto"
                repo.update_budget(g.cx, g.bid, title, currency, accent, -opening if neg else opening, theme)
                flash("Οι ρυθμίσεις αποθηκεύτηκαν ✓", "ok")
            except ValueError:
                flash("Μη έγκυρο αρχικό υπόλοιπο", "err")
            return redirect(url_for("settings"))
        return render_template("settings.html")

    # ---------- export ----------

    def export_month():
        month = request.args.get("m", "")
        return month if MONTH_RE.match(month) else ""

    @app.route("/export.csv")
    def export_csv():
        month = export_month()
        buf = io.StringIO()
        w = csv.writer(buf, delimiter=";")
        w.writerow(["Ημερομηνία", "Τύπος", "Κατηγορία", "Ποσό", "Σημείωση"])
        for r in repo.list_transactions(g.cx, g.bid, month=month or None):
            w.writerow([safe_cell(v) for v in (r["tdate"], "Έσοδο" if r["kind"] == "income" else "Έξοδο",
                                                r["cat_name"], fmt(r["amount_cents"], "").strip(), r["note"])])
        data = io.BytesIO(("﻿" + buf.getvalue()).encode("utf-8"))
        return send_file(data, mimetype="text/csv", as_attachment=True,
                         download_name=f"oikonomika_{month or 'all'}.csv")

    @app.route("/export.xlsx")
    def export_xlsx():
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        month = export_month()
        wb = Workbook()
        wb.properties.creator = "Oikonomika"
        wb.properties.lastModifiedBy = "Oikonomika"
        ws = wb.active
        ws.title = "Κινήσεις"
        ws.append(["Ημερομηνία", "Τύπος", "Κατηγορία", "Ποσό", "Σημείωση"])
        for r in repo.list_transactions(g.cx, g.bid, month=month or None):
            amt = r["amount_cents"] / 100
            ws.append([date.fromisoformat(r["tdate"]), "Έσοδο" if r["kind"] == "income" else "Έξοδο",
                       safe_cell(r["cat_name"]), amt if r["kind"] == "income" else -amt, safe_cell(r["note"])])
        for row in ws.iter_rows(min_row=2):
            row[0].number_format = "DD/MM/YYYY"
            row[3].number_format = '#,##0.00 "€";[Red]-#,##0.00 "€"'
        if month:
            s = repo.month_summary(g.cx, g.bid, month)
            ws2 = wb.create_sheet("Σύνοψη")
            ws2.append(["Κατηγορία", "Τύπος", "Πραγματικό", "Προϋπολογισμός", "Διαφορά"])
            for c in s["categories"]:
                ws2.append([safe_cell(c["name"]), "Έσοδο" if c["kind"] == "income" else "Έξοδο",
                            c["actual"] / 100, c["budget_cents"] / 100, (c["budget_cents"] - c["actual"]) / 100])
            ws2.append([])
            ws2.append(["Σύνολο εσόδων", "", s["income"] / 100, s["budget_income"] / 100])
            ws2.append(["Σύνολο εξόδων", "", s["expense"] / 100, s["budget_expense"] / 100])
            ws2.append(["Αποτέλεσμα", "", s["net"] / 100, s["budget_net"] / 100])
            for row in ws2.iter_rows(min_row=2):
                for cell in row[2:5]:
                    cell.number_format = '#,##0.00 "€";[Red]-#,##0.00 "€"'
        for sheet in wb.worksheets:
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="0F766E")
            for col, width in zip("ABCDE", (14, 10, 34, 16, 40)):
                sheet.column_dimensions[col].width = width
        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return send_file(out, as_attachment=True, download_name=f"oikonomika_{month or 'all'}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    return app
